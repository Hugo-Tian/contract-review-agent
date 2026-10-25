"""Generate entirely fictional parser fixtures, not real agreements."""
from pathlib import Path
import pymupdf
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'examples'
OUT.mkdir(exist_ok=True)
LINES = [
    '虚构示例：星桥求职辅导服务合同',
    '第一条 签约及收款主体：星桥示例教育有限公司。',
    '第二条 优惠成交价9800元。服务期2026年11月1日至2027年4月30日。',
    '第三条 简历修改2次，3个工作日内反馈。一对一辅导4次，每次60分钟。',
    '第四条 建群并开通账号即视为课程包全部交付。',
    '第五条 退出时按分项标价扣除已提供服务，另扣成交价20%的违约金。',
    '第六条 退款需提交学习记录，审核及到账期限另行通知。',
    '第七条 不保证正式录用，面试通过不等于录用。',
    '第八条 岗位数量和地区见附件一；本示例未附附件一。',
    '第九条 争议由双方协商处理。所有名称、交易和合同均为虚构。',
]
TABLE = [('服务项目', '分项标价', '交付标准'), ('课程包', '6000元', '建群并开通账号'), ('一对一辅导', '8000元', '4次，每次60分钟')]
word = Document()
for line in LINES:
    word.add_paragraph(line)
table = word.add_table(rows=0, cols=3)
table.style = 'Table Grid'
for row in TABLE:
    for cell, value in zip(table.add_row().cells, row):
        cell.text = value
word.save(OUT / 'fictional-contract.docx')
pdf = pymupdf.open()
page = pdf.new_page()
page.insert_text((45, 55), '\n\n'.join(LINES + [' | '.join(r) for r in TABLE]), fontname='china-s', fontsize=11)
pdf.save(OUT / 'fictional-contract.pdf')
pdf.close()
(OUT / 'fictional-contract.txt').write_text('\n'.join(LINES + [' | '.join(r) for r in TABLE]), encoding='utf-8')
print('Created fictional PDF, DOCX and TXT examples.')
