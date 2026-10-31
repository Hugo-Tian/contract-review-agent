"""Isolated parser worker. No network, filenames or original files are persisted."""
import io
import json
import sys
import uuid
import zipfile
from defusedxml import ElementTree
from .schemas import Block, Document
from .profiles import current

MAX_FILE = 10 * 1024 * 1024
MAX_PAGES = 80
MAX_TEXT = 250000


def parse_file(data: bytes, name: str, ocr_mode: str = 'auto') -> Document:
    if len(data) > MAX_FILE:
        raise ValueError('文件超过 10 MB 限制。')
    doc_id = uuid.uuid4().hex
    blocks: list[Block] = []
    warnings: list[str] = []

    def add(text: str, location: str, source: str = 'file'):
        if text.strip():
            blocks.append(Block(id=f'{doc_id}-b{len(blocks)+1:04}', location=location,
                                text=text.strip(), source=source, status='needs_review' if source == 'ocr' else 'extracted'))
            if sum(len(b.text) for b in blocks) > MAX_TEXT:
                raise ValueError('提取正文超过 250000 字符，未进行截断，请拆分文档。')

    if data.startswith(b'%PDF-') and name.lower().endswith('.pdf'):
        import pymupdf
        try:
            pdf = pymupdf.open(stream=data, filetype='pdf')
        except Exception:
            raise ValueError('PDF 损坏或无法打开。') from None
        with pdf:
            if pdf.needs_pass:
                raise ValueError('PDF 已加密，请先解密后上传。')
            if len(pdf) > MAX_PAGES:
                raise ValueError('PDF 超过 80 页限制，请拆分文档。')
            success = 0
            ocr_pages = 0
            for n, page in enumerate(pdf, 1):
                try:
                    text = page.get_text('text', sort=True).strip()
                except Exception:
                    warnings.append(f'第 {n} 页提取失败，未覆盖。')
                    continue
                source = 'file'
                if ocr_mode == 'force' or (ocr_mode == 'auto' and len(text) < 30):
                    if ocr_pages >= 20:
                        raise ValueError('一次最多 OCR 20 页，请拆分扫描件。')
                    from .ocr import recognize
                    scale = min(2, 2400 / max(page.rect.width, page.rect.height))
                    pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
                    recognized, low = recognize(pix.tobytes('png'))
                    ocr_pages += 1
                    if recognized.strip():
                        text, source = recognized, 'ocr'
                        warnings.append(f'第 {n} 页使用本机OCR；{low} 行识别把握较低。金额、日期、否定词必须核对原件。')
                if not text:
                    warnings.append(f'第 {n} 页没有识别到文字，未覆盖。')
                    continue
                success += 1
                add(text, f'PDF 第 {n} 页', source)
            coverage = f'PDF 共 {len(pdf)} 页，提取到文字 {success} 页；其中尝试OCR {ocr_pages} 页'
        warnings.append('PDF 表格、多栏、图片和签章可能无法完整还原，请逐页核对金额与对应服务。')
        complete = success == n
        kind = 'pdf'
    elif name.lower().endswith(('.png', '.jpg', '.jpeg', '.webp')):
        from .ocr import recognize
        try:
            text, low = recognize(data)
        except ValueError:
            raise
        except Exception:
            raise ValueError('图片无法识别，请上传清晰的 PNG / JPEG / WebP 图片。') from None
        add(text, '图片 第 1 页', 'ocr')
        kind, complete = 'image', bool(text.strip())
        coverage = '图片 1 张，本机 OCR 识别；不识别印章真伪或图片含义'
        warnings.append(f'OCR 有 {low} 行识别把握较低。请对照原图核对金额、日期、否定词及表格对应关系。')
    elif data.startswith(b'PK') and name.lower().endswith('.docx'):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                infos = archive.infolist()
                if len(infos) > 2000 or sum(i.file_size for i in infos) > 30 * 1024 * 1024:
                    raise ValueError('DOCX 解压体积或文件数量超出限制。')
                if 'word/document.xml' not in archive.namelist():
                    raise ValueError('文件不是有效 DOCX。')
                ElementTree.fromstring(archive.read('word/document.xml'))
            from docx import Document as WordDocument
            from docx.table import Table
            from docx.text.paragraph import Paragraph
            word = WordDocument(io.BytesIO(data))
            pno = tno = 0
            for item in word.iter_inner_content():
                if isinstance(item, Paragraph):
                    pno += 1
                    add(item.text, f'DOCX 正文段落 {pno}')
                elif isinstance(item, Table):
                    tno += 1
                    seen = set()
                    for rno, row in enumerate(item.rows, 1):
                        cells = []
                        for cno, cell in enumerate(row.cells, 1):
                            if cell._tc in seen:
                                cells.append(f'列{cno}：[合并单元格，见前一出现位置]')
                                continue
                            seen.add(cell._tc)
                            value = cell.text
                            if cell.tables:
                                warnings.append(f'表格 {tno} 含嵌套表格，未提取嵌套内容。')
                            cells.append(f'列{cno}：{value}')
                        add(' | '.join(cells), f'DOCX 表格 {tno} 第 {rno} 行')
            complete = not warnings
            coverage = f'DOCX 正文 {pno} 段、{tno} 个表格（不提供虚构页码）'
            warnings.append('DOCX 仅提取正文段落和顶层表格；页眉、页脚、文本框、批注、修订痕迹和图片未覆盖。自动条款编号可能需手动补充。')
            kind = 'docx'
        except ValueError:
            raise
        except Exception:
            raise ValueError('DOCX 损坏或无法解析。') from None
    else:
        raise ValueError('文件内容与类型不匹配，支持 PDF / DOCX / PNG / JPEG / WebP。')
    if not blocks:
        raise ValueError('没有可分析正文：未识别到文字，请上传更清晰的文件或手动粘贴文字。')
    if any(b.source == 'ocr' for b in blocks) and current().messages.get('ocr_warning'):
        warnings.append(current().messages['ocr_warning'])
    return Document(id=doc_id, name=name, kind=kind, blocks=blocks,
                    warnings=warnings, coverage=coverage, complete=complete)


if __name__ == '__main__':
    try:
        result = parse_file(sys.stdin.buffer.read(MAX_FILE + 1), sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else 'auto')
        sys.stdout.buffer.write(result.model_dump_json().encode('utf-8'))
    except Exception as exc:
        message = str(exc) if isinstance(exc, ValueError) else '文档解析失败。'
        sys.stdout.buffer.write(json.dumps({'error': message}, ensure_ascii=False).encode('utf-8'))
        sys.exit(1)
